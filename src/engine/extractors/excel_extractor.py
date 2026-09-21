# -*- coding: utf-8 -*-
"""
ExcelExtractor — BaseExtractor wrapper for .xls, .xlsx, .xlsm files.

Uses openpyxl for .xlsx/.xlsm and xlrd for .xls.
Returns register_row records compatible with ExtractJob persistence.
"""
from __future__ import annotations

import logging
import os
from typing import Any, Dict, List, Optional

from src.engine.extractors.base import BaseExtractor, ExtractionResult

logger = logging.getLogger(__name__)


class ExcelExtractor(BaseExtractor):
    """Extracts structured data from Excel files."""

    maturity = "VERIFIED"

    @property
    def id(self) -> str:
        return "excel-extractor-v1"

    @property
    def version(self) -> str:
        return "1.0.0"

    @property
    def supported_extensions(self) -> List[str]:
        return [".xls", ".xlsx", ".xlsm"]

    def extract(self, file_path: str, context: Optional[Dict[str, Any]] = None) -> ExtractionResult:
        context = context or {}
        source_path = os.path.abspath(file_path or "")
        file_name = os.path.basename(source_path)
        doc_id = context.get("doc_id", "unknown")
        ext = os.path.splitext(source_path)[1].lower()

        records: List[Dict[str, Any]] = []
        diagnostics: List[str] = []

        def update_stage(stage: str, msg: str = "", **extra: Any) -> None:
            cb = context.get("on_stage")
            if callable(cb):
                cb(stage, msg, **extra)

        try:
            if not os.path.exists(source_path):
                raise FileNotFoundError(f"File not found: {source_path}")

            update_stage("INITIALIZING", f"Loading Excel file: {ext}", source_path=source_path)

            if ext == ".xls":
                records, diagnostics = self._extract_xls(source_path, doc_id)
            elif ext in (".xlsx", ".xlsm"):
                records, diagnostics = self._extract_xlsx(source_path, doc_id)
            else:
                raise ValueError(f"Unsupported Excel extension: {ext}")

            update_stage("FINALIZING", f"Extracted {len(records)} row record(s)")

            return ExtractionResult(
                records=records,
                diagnostics=diagnostics,
                metadata={
                    "row_count": len(records),
                    "doc_id": doc_id,
                    "source_path": source_path,
                    "file_name": file_name,
                    "file_size": self._safe_size(source_path),
                },
                success=True,
            )

        except Exception as exc:
            logger.exception("[ExcelExtractor] Failed on %s", source_path)
            return ExtractionResult(
                success=False,
                diagnostics=[str(exc)],
                metadata={
                    "row_count": 0,
                    "doc_id": doc_id,
                    "source_path": source_path,
                    "file_name": file_name,
                    "file_size": self._safe_size(source_path),
                },
            )

    def _extract_xls(self, source_path: str, doc_id: str) -> tuple[list[dict], list[str]]:
        """Extract data from .xls using xlrd."""
        records: List[Dict[str, Any]] = []
        diagnostics: List[str] = []

        try:
            import xlrd
        except ImportError:
            diagnostics.append("xlrd is not installed. Cannot parse .xls files.")
            return records, diagnostics

        try:
            wb = xlrd.open_workbook(source_path)
            for sheet in wb.sheets():
                sheet_name = sheet.name
                for row_idx in range(sheet.nrows):
                    row_data = {}
                    for col_idx in range(sheet.ncols):
                        cell = sheet.cell(row_idx, col_idx)
                        value = cell.value
                        if value != "":
                            row_data[f"col_{col_idx}"] = value
                    if row_data:
                        records.append({
                            "type": "register_row",
                            "data": {
                                "sheet_name": sheet_name,
                                "row_index": row_idx + 1,
                                "content": row_data,
                            },
                            "provenance": {"sheet": sheet_name, "row": row_idx + 1},
                        })
            diagnostics.append(f"xlrd parsed {len(records)} row(s) from {len(wb.sheets())} sheet(s)")
        except Exception as e:
            diagnostics.append(f"xlrd failed: {e}")
            logger.exception("Failed to parse .xls with xlrd: %s", source_path)

        return records, diagnostics

    def _extract_xlsx(self, source_path: str, doc_id: str) -> tuple[list[dict], list[str]]:
        """Extract data from .xlsx/.xlsm using openpyxl."""
        records: List[Dict[str, Any]] = []
        diagnostics: List[str] = []

        try:
            import openpyxl
        except ImportError:
            diagnostics.append("openpyxl is not installed. Cannot parse .xlsx/.xlsm files.")
            return records, diagnostics

        try:
            wb = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    row_data = {}
                    for col_idx, value in enumerate(row, start=1):
                        if value is not None and value != "":
                            row_data[f"col_{col_idx}"] = value
                    if row_data:
                        records.append({
                            "type": "register_row",
                            "data": {
                                "sheet_name": sheet_name,
                                "row_index": row_idx,
                                "content": row_data,
                            },
                            "provenance": {"sheet": sheet_name, "row": row_idx},
                        })
            diagnostics.append(f"openpyxl parsed {len(records)} row(s) from {len(wb.sheetnames)} sheet(s)")
            wb.close()
        except Exception as e:
            diagnostics.append(f"openpyxl failed: {e}")
            logger.exception("Failed to parse .xlsx/.xlsm with openpyxl: %s", source_path)

        return records, diagnostics

    def _safe_size(self, path: str) -> int:
        try:
            return int(os.path.getsize(path))
        except Exception:
            return 0
